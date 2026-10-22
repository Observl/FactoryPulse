"""Auditable schema inference for heterogeneous industrial tables."""
from __future__ import annotations

import csv
import io
import json
import math
import re
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from .field_profiles import EXTRA_FIELDS, EXTRA_UNITS, custom_key, is_custom, leakage_column, future_outcome_column

STANDARD_FIELDS = {
    "equipment_id": "设备编号", "timestamp": "采集时间", "vibration_g": "振动RMS（g）",
    "vibration_mm_s": "振动速度（mm/s）", "vibration_peak_g": "振动峰值（g）",
    "vibration_kurtosis": "振动峭度", "process_temperature_k": "设备温度（K）",
    "rotational_speed_rpm": "转速（rpm）", "torque_nm": "扭矩（N·m）",
    "tool_wear_min": "刀具使用时间（min）", "spindle_load_pct": "负载（%）",
    "failure": "故障标签", "status": "设备状态", "production_line": "产线",
    "oil_temperature_c":"油温（℃）","motor_current_a":"电机电流（A）",
    "tp2_bar":"压力TP2（bar）","tp3_bar":"压力TP3（bar）","reservoirs_bar":"储气罐压力（bar）",
}

STANDARD_UNITS={"process_temperature_k":"K","oil_temperature_c":"℃","motor_current_a":"A","tp2_bar":"bar","tp3_bar":"bar","reservoirs_bar":"bar","vibration_g":"g","vibration_mm_s":"mm/s","vibration_peak_g":"g","rotational_speed_rpm":"rpm","torque_nm":"N·m","tool_wear_min":"min","spindle_load_pct":"%"}
STANDARD_FIELDS.update(EXTRA_FIELDS)
STANDARD_UNITS.update(EXTRA_UNITS)
STANDARD_UNITS["timestamp"]="日期时间"

def csv_delimiter(header):
    try: return csv.Sniffer().sniff(header,delimiters=",;\t|").delimiter
    except csv.Error: return ","


def iter_table_file(path: Path, original_filename: str):
    """Stream every record without imposing an artificial row/file-size limit."""
    suffix=Path(original_filename).suffix.lower()
    if suffix==".xlsx":
        workbook=load_workbook(path,read_only=True,data_only=True)
        try:
            iterator=iter(workbook.active.iter_rows(values_only=True))
            headers=[str(value or "").strip() for value in next(iterator)]
            if not headers or any(not key for key in headers) or len(set(headers))!=len(headers): raise ValueError("表头不能为空或重复")
            for values in iterator:
                if all(value in (None,"") for value in values): continue
                yield {key:value.isoformat(sep=" ") if isinstance(value,datetime) else value for key,value in zip(headers,values)}
        finally: workbook.close()
    elif suffix==".csv":
        import codecs
        with path.open("rb") as handle: prefix=handle.read(65536)
        encoding=None
        for candidate in ("utf-8-sig","gb18030"):
            try: codecs.getincrementaldecoder(candidate)().decode(prefix,final=False);encoding=candidate;break
            except UnicodeDecodeError: pass
        if not encoding: raise ValueError("CSV编码无法识别")
        with path.open("r",encoding=encoding,newline="") as handle:
            header=handle.readline();handle.seek(0)
            reader=csv.DictReader(handle,delimiter=csv_delimiter(header))
            headers=reader.fieldnames or []
            if not headers or any(not key.strip() for key in headers) or len(set(headers))!=len(headers): raise ValueError("表头不能为空或重复")
            for row in reader:
                if None in row: raise ValueError("数据列数多于表头，请检查分隔符和引号")
                yield row
    else: raise ValueError("仅支持CSV和XLSX")

ALIASES = {
    "equipment_type":[r"equipment.?type",r"machine.?type",r"设备类型",r"机型"],
    "air_temperature_k":[r"ambient.*temp",r"air.*temp",r"环境温度"],
    "pressure_bar":[r"^pressure(_bar)?$",r"工作压力",r"液压压力",r"出口压力",r"排气压力"],
    "flow_m3_h":[r"flow",r"流量"],"voltage_v":[r"voltage",r"电压"],
    "power_kw":[r"power.*kw",r"功率"],"frequency_hz":[r"frequency",r"频率"],
    "humidity_pct":[r"humidity",r"湿度"],"operating_hours":[r"operating.?hours",r"运行时长",r"累计运行"],
    "cycle_count":[r"cycle.?count",r"循环次数"],"displacement_mm":[r"displacement",r"位移"],
    "oil_level_pct":[r"oil.?level",r"油位"],"particle_count":[r"particle",r"颗粒"],
    "oil_temperature_c":[r"oil.?temperature.?c?",r"油温",r"润滑油温度"],
    "motor_current_a":[r"motor.?current.?a?",r"电机电流"],
    "tp2_bar":[r"tp2(_bar)?"],"tp3_bar":[r"tp3(_bar)?"],"reservoirs_bar":[r"reservoirs(_bar)?",r"储气罐压力"],
    "equipment_id": [r"equipment.?id", r"machine.?id", r"device.?id", r"asset.?id", r"设备.*(编号|id)", r"机台.*(编号|id)", r"序列号", r"serial"],
    "timestamp": [r"timestamp", r"date.?time", r"record.?time", r"采集时间", r"记录时间", r"日期时间", r"时间戳", r"^time$"],
    "vibration_g": [r"vibration.*rms.*g", r"accel.*rms", r"振动.*rms.*g", r"加速度.*rms"],
    "vibration_mm_s": [r"vibration.*mm.?s", r"velocity.*rms", r"振动速度", r"振动.*mm"],
    "vibration_peak_g": [r"vibration.*peak", r"accel.*peak", r"振动峰值", r"加速度峰值"],
    "vibration_kurtosis": [r"kurtosis", r"峭度"],
    "process_temperature_k": [r"process.*temp", r"bearing.*temp", r"spindle.*temp", r"设备温度", r"轴承温度", r"主轴温度", r"温度", r"temperature"],
    "rotational_speed_rpm": [r"rotational.*speed", r"speed.*rpm", r"转速", r"rpm"],
    "torque_nm": [r"torque", r"扭矩"], "tool_wear_min": [r"tool.*wear", r"刀具.*(时间|磨损)", r"wear.*min"],
    "spindle_load_pct": [r"spindle.*load", r"load.*%", r"主轴负载", r"负载率"],
    "failure": [r"future.?failure", r"failure", r"fault", r"故障", r"失效"],
    "status": [r"status", r"state", r"设备状态", r"运行状态"],
    "production_line": [r"production.?line", r"line.?id", r"产线", r"生产线"],
}


def parse_table(filename: str, raw: bytes, limit: int = 200_000) -> tuple[str, list[dict[str, Any]]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator)]
        rows = []
        for values in iterator:
            if all(value in (None, "") for value in values): continue
            rows.append({key: value.isoformat(sep=" ") if isinstance(value, datetime) else value for key, value in zip(headers, values) if key})
            if len(rows) >= limit: break
        workbook.close(); return sheet.title, rows
    if suffix == ".csv":
        text = None
        for encoding in ("utf-8-sig", "gb18030"):
            try: text = raw.decode(encoding); break
            except UnicodeDecodeError: pass
        if text is None: raise ValueError("CSV编码无法识别，请使用UTF-8或GB18030")
        reader = csv.DictReader(io.StringIO(text),delimiter=csv_delimiter(text.splitlines()[0] if text else ""))
        rows = [dict(row) for _, row in zip(range(limit), reader)]
        return "CSV", rows
    raise ValueError("仅支持 .xlsx 或 .csv 文件")


def parse_table_path(path: Path, original_filename: str, sample_only: bool = False) -> tuple[str, list[dict[str, Any]]]:
    """Read one example row for mapping, or the complete file after confirmation."""
    suffix = Path(original_filename).suffix.lower()
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator)]
        rows = []
        for values in iterator:
            if all(value in (None, "") for value in values):
                continue
            rows.append({key: value.isoformat(sep=" ") if isinstance(value, datetime) else value for key, value in zip(headers, values) if key})
            if sample_only:
                break
        workbook.close()
        return sheet.title, rows
    if suffix == ".csv":
        encoding = None
        with path.open("rb") as probe:
            prefix = probe.read(64 * 1024)
        for candidate in ("utf-8-sig", "gb18030"):
            try:
                prefix.decode(candidate); encoding = candidate; break
            except UnicodeDecodeError:
                pass
        if encoding is None:
            raise ValueError("CSV编码无法识别，请使用UTF-8或GB18030")
        with path.open("r", encoding=encoding, newline="") as handle:
            header=handle.readline();handle.seek(0)
            reader = csv.DictReader(handle,delimiter=csv_delimiter(header))
            if not reader.fieldnames:
                raise ValueError("CSV缺少表头")
            rows = []
            for row in reader:
                rows.append(dict(row))
                if sample_only:
                    break
        return "CSV", rows
    raise ValueError("仅支持 .xlsx 或 .csv 文件")


def infer_mapping(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not rows: raise ValueError("数据表没有有效记录")
    output = []
    for column in rows[0]:
        normalized = re.sub(r"[\s_\-()/（）%]+", "", column.lower())
        best = ("", 0.0, "未识别")
        if column in STANDARD_FIELDS: best=(column,.99,"标准字段名称一致")
        for target, patterns in ALIASES.items():
            for pattern in patterns:
                if re.search(pattern, column.lower(), re.I) or re.search(pattern, normalized, re.I):
                    score = .97 if re.fullmatch(pattern, column.lower(), re.I) else .84
                    if score > best[1]: best = (target, score, f"字段名匹配：{pattern}")
        samples = [rows[0].get(column)] if rows[0].get(column) not in (None, "") else []
        custom=custom_key(column)
        if future_outcome_column(column):
            best=("",0.,"结果或未来标签不能作为当前故障状态，请提供实际发生时的故障标记")
        if not best[0] and samples and not leakage_column(column):
            try:
                if math.isfinite(float(samples[0])): best=(custom,.55,"数值型扩展指标，需确认含义及单位")
            except (ValueError,TypeError): pass
        output.append({"source_column": column, "target_field": best[0], "custom_target":custom,"target_label": STANDARD_FIELDS.get(best[0], column if is_custom(best[0]) else "忽略"), "confidence": best[1], "reason": best[2], "samples": samples,"source_unit":"standard","use_for_prediction":best[0] not in {"","equipment_id","timestamp","failure","status","production_line","equipment_type"}})
    used: dict[str, dict[str, Any]] = {}
    for item in sorted(output, key=lambda value: value["confidence"], reverse=True):
        target = item["target_field"]
        if target and target in used:
            if target not in {"equipment_id","timestamp","failure","status","production_line","equipment_type"}:
                item.update(target_field=item["custom_target"],target_label=item["source_column"],confidence=.55,reason="同类传感器的另一测点，作为独立扩展指标，请确认单位")
            else: item.update({"target_field": "", "target_label": "忽略", "confidence": 0.0, "reason": f"与 {used[target]['source_column']} 重复"})
        elif target: used[target] = item
    return output


def normalize_rows(rows: list[dict[str, Any]], mapping: dict[str, str], default_equipment_id: str) -> list[dict[str, Any]]:
    if any(target=="failure" and future_outcome_column(column) for column,target in mapping.items()):
        raise ValueError("未来结果或寿命字段不能映射为当前故障状态，请提供实际发生时的故障标记")
    normalized=[]; errors=[]; seen=set()
    numeric=(set(STANDARD_FIELDS)|{field for field in mapping.values() if is_custom(field)})-{"equipment_id","timestamp","status","production_line","equipment_type"}
    for index,source in enumerate(rows):
        row={target:source.get(column) for column,target in mapping.items() if target}
        row["equipment_id"]=str(row.get("equipment_id") or default_equipment_id).strip()
        try:
            if not row["equipment_id"]: raise ValueError("设备编号为空，请补充编号或明确指定单设备文件的固定编号")
            text=str(row.get("timestamp") or "").strip()
            try: moment=datetime.fromisoformat(text)
            except ValueError:
                moment=None
                for pattern in ("%Y/%m/%d %H:%M:%S","%Y/%m/%d %H:%M","%Y/%m/%d","%Y年%m月%d日 %H:%M:%S","%Y%m%d%H%M%S"):
                    try: moment=datetime.strptime(text,pattern);break
                    except ValueError: pass
                if moment is None: raise ValueError("采集时间格式无法识别，请使用含年份的日期时间；时间戳需指定秒或毫秒单位")
            # Store a consistent UTC-naive representation when timezone is supplied.
            if moment.tzinfo:
                from datetime import timezone
                moment=moment.astimezone(timezone.utc).replace(tzinfo=None)
            row["timestamp"]=moment.isoformat(sep=" ")
            key=(row["equipment_id"],row["timestamp"])
            if key in seen: raise ValueError("设备与采集时间重复")
            seen.add(key)
            for field in numeric:
                if field in row and row[field] not in (None,""):
                    if field=="failure" and isinstance(row[field],str):
                        text=row[field].strip().lower()
                        row[field]={"true":1,"false":0,"yes":1,"no":0,"故障":1,"正常":0,"是":1,"否":0}.get(text,row[field])
                    value=float(row[field])
                    if not math.isfinite(value): raise ValueError(f"{field} 不是有限数值")
                    if field=="failure" and value not in (0,1): raise ValueError("故障标签只能是0或1")
                    row[field]=value
                elif field in row: row[field]=None
            row["reading_id"]=f"IMPORT-{row['equipment_id']}-{index:08d}"
            row["data_origin"]="AI_ASSISTED_USER_IMPORT"
            normalized.append(row)
        except (ValueError,TypeError) as exc:
            if len(errors)<20: errors.append(f"第 {index+2} 行：{exc}")
    if errors: raise ValueError("数据校验未通过，未写入任何记录。"+"；".join(errors))
    if not normalized: raise ValueError("没有可导入的监测记录")
    normalized.sort(key=lambda r:(r["timestamp"],r["equipment_id"]))
    grouped={}
    for row in normalized: grouped.setdefault(row["equipment_id"],[]).append(row)
    for history in grouped.values():
        end=datetime.fromisoformat(history[-1]["timestamp"])
        onsets=[]; previous=None
        for row in history:
            if row.get("failure")==1 and previous==0: onsets.append(datetime.fromisoformat(row["timestamp"]))
            previous=row.get("failure")
        complete_labels=all(row.get("failure") in (0,1) for row in history)
        for row in history:
            moment=datetime.fromisoformat(row["timestamp"])
            row["prediction_horizon_hours"]=168
            row["future_failure"]=None if not complete_labels or row.get("failure")==1 or moment+timedelta(hours=168)>end else int(any(moment<onset<=moment+timedelta(hours=168) for onset in onsets))
            row["label_status"]="observed" if row["future_failure"] is not None else "unknown_or_already_failed"
    return normalized
