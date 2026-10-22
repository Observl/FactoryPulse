import io
import os
import asyncio
from pathlib import Path
import httpx
from openpyxl import Workbook
os.environ.setdefault("FACTORYPULSE_DATA_DIR",str(Path(__file__).parent/".runtime"))
from backend.main import app, TABLES

class ASGIClient:
    def request(self,method,path,**kwargs):
        async def run():
            transport=httpx.ASGITransport(app=app)
            async with httpx.AsyncClient(transport=transport,base_url="http://test") as client:
                return await client.request(method,path,**kwargs)
        return asyncio.run(run())
    def get(self,path,**kwargs): return self.request("GET",path,**kwargs)
    def post(self,path,**kwargs): return self.request("POST",path,**kwargs)
client=ASGIClient()
def workbook():
    w=Workbook();w.remove(w.active)
    for table,(sheet,required) in TABLES.items():
        required=[*required,"skills"] if table=="staff" else required
        s=w.create_sheet(sheet)
        s.append(required)
        record={name: "x" for name in required}
        record.update({"equipment_id":"CNC-001","equipment_name":"测试数控机床","reading_id":"R-1","timestamp":"2026-10-15 08:00:00","scenario_id":"S-1","risk_score":0.8,"maintenance_budget_cny":2000,"risk_threshold":0.65,"available_staff_count":2,"max_downtime_h":8,"inspection_cost_cny":100,"repair_cost_cny":500,"inspection_duration_h":1,"repair_duration_h":2,"required_staff_count":1,"assumed_failure_loss_cny":5000,"staff_id":"ST-1","skills":"机修","available_hours":8,"order_id":"O-1","maintenance_id":"M-1","rotational_speed_rpm":1300,"torque_nm":25,"air_temperature_k":298,"process_temperature_k":310,"tool_wear_min":80})
        s.append([record[k] for k in required])
    b=io.BytesIO();w.save(b);return b.getvalue()
def test_health():
    assert client.get("/api/health").status_code==200
def test_import_and_simulate():
    r=client.post("/api/import",files={"file":("test.xlsx",workbook(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code==200,r.text
    assert r.json()["counts"]["equipment"]==1
    r=client.get("/api/overview")
    assert r.status_code==200
    assert r.json()["metrics"]["total"]==1
    assert r.json()["devices"][0]["risk_score"]==.8
    r=client.post("/api/simulate",json={"scenario_id":"S-1","strategy":"optimized"})
    assert r.status_code==200,r.text
    assert r.json()["summary"]["maintenance_cost"]<=2000
    assert r.json()["summary"]["maintenance_hours"]<=8
def test_reject_non_xlsx():
    r=client.post("/api/import",files={"file":("bad.txt",b"hello","text/plain")})
    assert r.status_code==400
def test_developer_settings_do_not_expose_key(monkeypatch):
    monkeypatch.setattr('backend.main.require_public_https', lambda value: None)
    monkeypatch.setattr('backend.main.credential_set', lambda value: True)
    r=client.post("/api/developer/settings",json={"provider":"custom","base_url":"https://example.org/v1","model":"example","api_key":"synthetic-test-secret"})
    assert r.status_code==200
    assert "synthetic-test-secret" not in r.text
    got=client.get("/api/developer/settings")
    assert got.json()["has_api_key"] is True
    assert "synthetic-test-secret" not in got.text

def test_all_strategy_comparison_and_csv():
    uploaded=client.post("/api/import", files={"file":("test.xlsx", workbook(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert uploaded.status_code==200,uploaded.text
    response=client.get("/api/strategies/compare?scenario_id=S-1")
    assert response.status_code==200,response.text
    results=response.json()["strategies"]
    assert {x["strategy"] for x in results}=={"fixed_interval","risk_threshold","optimized"}
    assert all(x["summary"]["maintenance_cost"]<=2000 for x in results)
    assert all(x["summary"]["maintenance_hours"]<=8 for x in results)
    csv_result=client.get("/api/export/maintenance.csv?scenario_id=S-1")
    assert csv_result.status_code==200
    assert "text/csv" in csv_result.headers["content-type"]
    assert "设备ID" in csv_result.text

def test_invalid_risk_is_rejected():
    from openpyxl import load_workbook
    b=io.BytesIO(workbook()); wb=load_workbook(b)
    ws=wb["08_风险评分演示"]
    risk_col=list(next(ws.values)).index("risk_score")+1
    ws.cell(row=2,column=risk_col).value=1.5
    bad=io.BytesIO();wb.save(bad)
    res=client.post("/api/import",files={"file":("invalid.xlsx",bad.getvalue(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert res.status_code==422
    assert "风险评分" in res.text

def test_reject_local_api_destination():
    response=client.post("/api/developer/settings",json={"provider":"custom","base_url":"https://127.0.0.1/v1","model":"demo"})
    assert response.status_code==422

def test_alert_schedule_approval_and_audit():
    uploaded=client.post("/api/import",files={"file":("test.xlsx",workbook(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert uploaded.status_code==200
    alert=client.post("/api/alerts/CNC-001/action",json={"action":"work_order","note":"现场复核后确认维修需求"})
    assert alert.status_code==200
    assert alert.json()["status"]=="已转工单"
    schedule=client.post("/api/schedule/generate",json={"scenario_id":"S-1","start_at":"2026-10-18T08:00:00"})
    assert schedule.status_code==200,schedule.text
    assert schedule.json()["human_approval_required"] is True
    orders=client.get("/api/work-orders").json()["records"]
    assert orders
    approved=client.post(f"/api/work-orders/{orders[0]['work_order_id']}/approval",json={"decision":"approve","approver":"设备总管","note":"批准窗口","version":orders[0]["version"],"accept_assumptions":True})
    assert approved.status_code==200
    assert approved.json()["status"]=="已批准"
    audit=client.get("/api/audit").json()["records"]
    assert any(item["event_type"]=="work_order_approval" for item in audit)

def test_quality_and_evaluation_expose_evidence_state():
    assert client.get("/api/data-quality").status_code==200
    evaluation=client.get("/api/model/evaluation")
    assert evaluation.status_code==200
    assert "status" in evaluation.json()

def test_intelligent_field_mapping_requires_confirmation(monkeypatch):
    monkeypatch.setattr("backend.main._api_key", "")
    content="机台编号,采集时间,振动RMS(g),轴承温度,故障状态\nCNC-A,2026-10-01 08:00:00,0.12,310,0\nCNC-A,2026-10-01 09:00:00,0.18,312,1\n".encode("utf-8-sig")
    analyzed=client.post("/api/import/intelligent/analyze",files={"file":("company.csv",content,"text/csv")})
    assert analyzed.status_code==200,analyzed.text
    body=analyzed.json()
    proposed={item["source_column"]:item["target_field"] for item in body["mapping"]}
    assert proposed["机台编号"]=="equipment_id"
    assert proposed["采集时间"]=="timestamp"
    confirmed=client.post("/api/import/intelligent/confirm",json={"upload_id":body["upload_id"],"mapping":[{"source_column":column,"target_field":target} for column,target in proposed.items()]})
    assert confirmed.status_code==200,confirmed.text
    assert confirmed.json()["device_count"]==1
    assert confirmed.json()["record_count"]==2
    assert client.get("/api/overview").json()["sync"]["data_origin"]=="AI_ASSISTED_USER_IMPORT"
